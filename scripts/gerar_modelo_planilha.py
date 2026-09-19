"""Gera uma planilha-modelo XLSX sanitizada para versionamento."""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook as OpenpyxlWorkbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from preparar_insumos import HEADERS, HUMAN, ROOT, Workbook, column_name


OUTPUT = ROOT / "exemplos" / "controle_pipeline_flow.modelo.xlsx"
OWNERS = {
    name: (
        "VOCÊ/IA"
        if name in {"cta_destino", "cta_palavra"}
        else "VOCÊ"
        if name in HUMAN
        else "SISTEMA"
    )
    for name in HEADERS
}
DESCRIPTIONS = {
    "classifica": "sim processa a linha; não mantém apenas como referência.",
    "arquivo": "Imagem-base existente usada no clipe.",
    "produto_id": "Identificador estável do produto.",
    "producao_id": "Identificador que agrupa os clipes da produção.",
    "ordem": "Posição do clipe na produção.",
    "papel_na_producao": "abertura, principal ou cta.",
    "link_produto": "Link factual opcional; não é referência visual.",
    "usar_com": "Referências adicionais do mesmo clipe, separadas por ponto e vírgula.",
    "tipo_referencia": "Papel visual do arquivo principal.",
    "material_existente": "Ativo pronto que pode ser adaptado ou reutilizado.",
    "uso_material": "adaptar, reutilizar, referência ou automático.",
    "instrucao": "O que preservar, alterar e não copiar.",
    "status": "Estado geral preenchido pelo sistema.",
    "classificacao": "Categoria criativa definida pela classificação.",
    "fluxo": "Fluxo técnico selecionado no plano.",
    "plano_arquivo": "Caminho do plano técnico importado.",
    "imagem_status": "Estado da imagem ou frame.",
    "imagem_arquivo": "Caminho da imagem vinculada à revisão.",
    "aprovacao": "Decisão humana que libera exclusivamente o vídeo.",
    "video_status": "Estado da geração ou reutilização do vídeo.",
    "video_arquivo": "Caminho do vídeo concluído ou reutilizado.",
    "erro": "Mensagem operacional registrada pelo sistema.",
    "atualizado_em": "Data e hora da última atualização.",
    "fala_audio": "Fala humana opcional a preservar no planejamento.",
    "gerar_carrossel": "sim autoriza somente o card 9:16.",
    "cta_destino": "Destino autorizado da chamada no card CTA.",
    "cta_palavra": "Palavra-chave do card CTA.",
    "texto_carrossel": "Título criado pela IA, máximo de 52 caracteres.",
    "subtexto_carrossel": "Destaque criado pela IA, máximo de 28 caracteres.",
    "carrossel_status": "Estado da renderização local do card.",
    "carrossel_arquivo": "Caminho do card renderizado.",
}


def add_validation(sheet, target: str, values: list[str], strict: bool = True) -> None:
    formula = '"' + ",".join(values) + '"'
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=strict,
        showInputMessage=True,
    )
    validation.promptTitle = "Pipeline Flow"
    validation.prompt = "Escolha uma opção ou deixe vazio quando permitido."
    sheet.add_data_validation(validation)
    validation.add(target)


def style_header(sheet, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[row]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 34


def configure_control(sheet) -> None:
    sheet.title = "Controle"
    for index, name in enumerate(HEADERS, 1):
        sheet.cell(1, index, name)
        sheet.cell(2, index, OWNERS[name])
        sheet.cell(3, index, DESCRIPTIONS[name])
        sheet.column_dimensions[column_name(index - 1)].width = 20
    style_header(sheet)
    for row in (2, 3):
        for cell in sheet[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 42 if row == 3 else 22
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A1:{column_name(len(HEADERS) - 1)}500"
    add_validation(sheet, "A4:A500", ["sim", "não"])
    add_validation(sheet, "F4:F500", ["abertura", "principal", "cta"])
    add_validation(
        sheet,
        "I4:I500",
        ["base_edicao", "produto", "inspiração", "detalhe", "ambiente", "outro"],
    )
    add_validation(sheet, "K4:K500", ["adaptar", "reutilizar", "referência", "automático"])
    add_validation(sheet, "S4:S500", ["aprovada", "rejeitada"])
    add_validation(sheet, "Y4:Y500", ["sim", "não"])
    add_validation(
        sheet,
        "Z4:Z500",
        ["grupo de achadinhos", "direct", "link da bio", "comentários", "WhatsApp"],
        strict=False,
    )
    add_validation(
        sheet,
        "AA4:AA500",
        ["EXEMPLO", "CASA", "ORGANIZAÇÃO", "ACHADINHOS"],
        strict=False,
    )
    green = PatternFill("solid", fgColor="E2F0D9")
    red = PatternFill("solid", fgColor="FCE4D6")
    sheet.conditional_formatting.add(
        "S4:S500", FormulaRule(formula=['S4="aprovada"'], fill=green)
    )
    sheet.conditional_formatting.add(
        "S4:S500", FormulaRule(formula=['S4="rejeitada"'], fill=red)
    )


def add_guide(workbook) -> None:
    sheet = workbook.create_sheet("Guia")
    sheet.append(["campo", "preenchido por", "descrição"])
    for name in HEADERS:
        sheet.append([name, OWNERS[name], DESCRIPTIONS[name]])
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 90
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_references(workbook) -> None:
    sheet = workbook.create_sheet("Referencias_Identidade")
    sheet.append(["tipo", "arquivo fictício", "observação"])
    sheet.append(["identidade_rosto", "referencias/rosto_exemplo.jpg", "Use somente com autorização."])
    sheet.append(["identidade_corpo", "referencias/corpo_exemplo.png", "Inclua apenas se o corpo aparecer."])
    sheet.append(["identidade_mao", "referencias/mao_exemplo.jpg", "Inclua somente quando necessária."])
    style_header(sheet)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 52


def add_example(workbook) -> None:
    sheet = workbook.create_sheet("Exemplo_Producao")
    sheet.append(HEADERS)
    examples = [
        {
            "classifica": "não",
            "arquivo": "exemplo_abertura.jpg",
            "produto_id": "PRODUTO_EXEMPLO",
            "producao_id": "PRODUCAO_EXEMPLO_01",
            "ordem": "1",
            "papel_na_producao": "abertura",
            "tipo_referencia": "base_edicao",
            "instrucao": "Preservar a cena-base e expandir as bordas para 9:16.",
            "gerar_carrossel": "sim",
        },
        {
            "classifica": "não",
            "arquivo": "exemplo_produto.jpg",
            "produto_id": "PRODUTO_EXEMPLO",
            "producao_id": "PRODUCAO_EXEMPLO_01",
            "ordem": "2",
            "papel_na_producao": "principal",
            "tipo_referencia": "base_edicao",
            "instrucao": "Preservar produto e cenário; alterar apenas o enquadramento.",
        },
        {
            "classifica": "não",
            "arquivo": "exemplo_cta.jpg",
            "produto_id": "PRODUTO_EXEMPLO",
            "producao_id": "PRODUCAO_EXEMPLO_01",
            "ordem": "3",
            "papel_na_producao": "cta",
            "tipo_referencia": "base_edicao",
            "instrucao": "Preservar a cena e adaptar para o encerramento.",
            "gerar_carrossel": "sim",
            "cta_destino": "salvar; compartilhar",
            "cta_palavra": "EXEMPLO",
        },
    ]
    for example in examples:
        sheet.append([example.get(name, "") for name in HEADERS])
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{column_name(len(HEADERS) - 1)}4"
    for index in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[column_name(index - 1)].width = 20


def normalize_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        parts = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in sorted(parts):
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-09-19T00:00:00Z\g<2>",
                    content,
                )
            info = zipfile.ZipInfo(name, (2026, 9, 19, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, content)
    temporary.replace(path)


def generate(output: Path) -> Path:
    workbook = OpenpyxlWorkbook()
    workbook.properties.creator = "Pipeline Flow"
    workbook.properties.title = "Modelo sanitizado do controle do Pipeline Flow"
    workbook.properties.created = datetime(2026, 9, 19)
    workbook.properties.modified = datetime(2026, 9, 19)
    configure_control(workbook.active)
    add_guide(workbook)
    add_references(workbook)
    add_example(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    normalize_zip(output)
    load_workbook(output, read_only=True).close()
    Workbook(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--saida", type=Path, default=OUTPUT)
    args = parser.parse_args()
    print(generate(args.saida.resolve()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
