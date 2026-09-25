import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook

from pipeline_flow.config import load_config
from pipeline_flow.services.preparar_insumos import HEADERS, signature
from pipeline_flow.web.app import _is_loopback, create_app
from pipeline_flow.web.queries import PipelineReadModel


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create_sheet(path, plan_path):
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = 'Controle'
    sheet.append(HEADERS)
    sheet.append([''] * len(HEADERS))
    sheet.append([''] * len(HEADERS))
    values = {
        'classifica': 'sim',
        'produto_id': 'PRODUTO_01',
        'producao_id': 'PROD_01',
        'ordem': '1',
        'papel_na_producao': 'principal',
        'plano_arquivo': str(plan_path),
        'status': 'aguardando_aprovacao',
        'imagem_status': 'gerada',
    }
    sheet.append([values.get(name, '') for name in HEADERS])
    book.save(path)


class WebReadOnlyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.output = self.root / 'preparados'
        self.delivery = self.root / 'entregas_flow'
        staging = self.output / 'flow' / 'PROD_01' / 'staging'
        clip_dir = staging / 'CLIP_01'
        clip_dir.mkdir(parents=True)
        self.delivery.mkdir()
        plan = {
            'id_clipe': 'CLIP_01',
            'producao_id': 'PROD_01',
            'produto_id': 'PRODUTO_01',
            'ordem': 1,
            'status': 'pronto',
            'pipeline': {
                'aspecto': '9:16',
                'gerar_imagem': False,
                'metodo_imagem': None,
                'gerar_video': False,
                'metodo_video': None,
                'duracao_video_s': None,
                'aprovacao_necessaria': False,
            },
        }
        summary = {
            'producao_id': 'PROD_01',
            'clipes': [{'id_clipe': 'CLIP_01', 'ordem': 1}],
        }
        response = {
            'pacote_sha256': 'b' * 64,
            'plano_producao': summary,
            'clipes': [{'plano': plan}],
        }
        revision = staging.with_name(signature(response)[:16])
        staging.rename(revision)
        self.revision = revision
        self.clip_dir = revision / 'CLIP_01'
        (revision / 'plano_producao.json').write_text(
            json.dumps(summary), encoding='utf-8'
        )
        (revision / 'resposta_ia.json').write_text(
            json.dumps(response), encoding='utf-8'
        )
        (self.clip_dir / 'plano_clipe.json').write_text(
            json.dumps(plan), encoding='utf-8'
        )
        inputs = {'pacote_sha256': 'b' * 64, 'referencias': []}
        (self.clip_dir / 'insumos_flow.json').write_text(
            json.dumps(inputs), encoding='utf-8'
        )
        fingerprint = signature({
            'plano': plan,
            'insumos': inputs,
            'prompts': {},
        })
        (self.clip_dir / 'execucao.json').write_text(
            json.dumps({'schema_version': 2, 'fingerprint': fingerprint}),
            encoding='utf-8',
        )
        (self.delivery / 'indice_entregas.json').write_text('[]', encoding='utf-8')
        self.preview = self.output / 'preview.png'
        self.preview.write_bytes(b'local preview')
        self.sheet = self.root / 'controle.xlsx'
        create_sheet(self.sheet, self.clip_dir / 'plano_clipe.json')
        base = load_config(self.root, environ={})
        self.config = replace(
            base,
            spreadsheet=self.sheet,
            output_dir=self.output,
            delivery_dir=self.delivery,
            web_host='127.0.0.1',
        )
        self.files = [
            self.sheet,
            revision / 'plano_producao.json',
            revision / 'resposta_ia.json',
            self.clip_dir / 'plano_clipe.json',
            self.clip_dir / 'insumos_flow.json',
            self.clip_dir / 'execucao.json',
            self.delivery / 'indice_entregas.json',
            self.preview,
        ]

    def hashes(self):
        return {path: file_hash(path) for path in self.files}

    def test_read_model_lists_active_production_without_writes(self):
        before = self.hashes()

        dashboard = PipelineReadModel(self.config).dashboard()

        self.assertEqual(dashboard['resumo']['producoes'], 1)
        self.assertEqual(dashboard['resumo']['clipes'], 1)
        self.assertEqual(dashboard['producoes'][0]['clipes'][0]['id'], 'CLIP_01')
        self.assertEqual(before, self.hashes())

    def test_api_exposes_only_read_routes_and_preserves_files(self):
        before = self.hashes()
        client = TestClient(create_app(self.config))

        self.assertEqual(client.get('/api/health').status_code, 200)
        self.assertEqual(client.get('/api/dashboard').status_code, 200)
        self.assertEqual(client.get('/api/productions').status_code, 200)
        self.assertEqual(client.get('/api/productions/PROD_01').status_code, 200)
        self.assertEqual(client.get('/api/clips').status_code, 200)
        self.assertEqual(client.get('/api/deliveries').status_code, 200)
        self.assertEqual(client.get('/api/logs').status_code, 200)
        self.assertEqual(client.post('/api/dashboard').status_code, 405)
        self.assertEqual(before, self.hashes())

    def test_html_and_static_assets_are_local_and_hardened(self):
        client = TestClient(create_app(self.config))

        page = client.get('/')
        css = client.get('/static/app.css')
        script = client.get('/static/app.js')

        self.assertEqual(page.status_code, 200)
        self.assertIn("id='kpi-grid'", page.text)
        self.assertEqual(css.status_code, 200)
        self.assertEqual(script.status_code, 200)
        self.assertIn('operations', page.text)
        self.assertIn('/api/operations', script.text)
        self.assertIn('review-grid', page.text)
        self.assertIn('/api/operations/review-image', script.text)
        self.assertIn('image_sha256', script.text)
        self.assertIn('carousel-list', page.text)
        self.assertIn('/api/operations/generate-carousel', script.text)
        self.assertIn('video-list', page.text)
        self.assertIn('/api/operations/generate-video', script.text)
        self.assertIn('/api/operations/executions', script.text)
        self.assertIn('GERAR VIDEO', page.text)
        self.assertIn('GERAR VIDEO', script.text)
        self.assertIn("id='assistant-tab'", page.text)
        self.assertIn("id='assistant-command'", page.text)
        self.assertIn('python scripts\\\\rodar_pipeline.py', script.text)
        self.assertIn('registrar-imagem', script.text)
        self.assertIn('--modelo-video omni-flash', script.text)
        self.assertIn('navigator.clipboard', script.text)
        quote = chr(39)
        self.assertIn('role=' + quote + 'tablist' + quote, page.text)
        self.assertEqual(page.text.count('role=' + quote + 'tab' + quote), 9)
        self.assertIn('aria-selected=' + quote + 'true' + quote, page.text)
        self.assertEqual(page.text.count('role=' + quote + 'tabpanel' + quote), 9)
        self.assertIn('aria-busy=' + quote + 'true' + quote, page.text)
        self.assertIn('ArrowRight', script.text)
        self.assertIn('Home', script.text)
        self.assertIn('aria-selected', script.text)
        self.assertIn('aria-busy', script.text)
        self.assertIn('scrollIntoView', script.text)
        self.assertIn('aria-hidden', script.text)
        self.assertEqual(page.headers['cache-control'], 'no-store')
        self.assertIn("default-src 'self'", page.headers['content-security-policy'])

    def test_media_route_serves_only_allowlisted_local_media(self):
        client = TestClient(create_app(self.config))

        media = client.get('/media/preparados/preview.png')
        manifest = client.get(
            f'/media/preparados/flow/PROD_01/{self.revision.name}/resposta_ia.json'
        )
        outside = client.get('/media/desconhecido/preview.png')

        self.assertEqual(media.status_code, 200)
        self.assertEqual(media.content, b'local preview')
        self.assertEqual(manifest.status_code, 404)
        self.assertEqual(outside.status_code, 404)

    def test_read_model_creates_media_url_only_inside_allowed_roots(self):
        model = PipelineReadModel(self.config)

        self.assertEqual(model._media_url(self.preview), '/media/preparados/preview.png')
        self.assertIsNone(model._media_url(self.sheet))

    def test_missing_production_returns_404(self):
        client = TestClient(create_app(self.config))
        self.assertEqual(
            client.get('/api/productions/INEXISTENTE').status_code,
            404,
        )

    def test_server_accepts_only_loopback_hosts(self):
        self.assertTrue(_is_loopback('127.0.0.1'))
        self.assertTrue(_is_loopback('::1'))
        self.assertTrue(_is_loopback('localhost'))
        self.assertFalse(_is_loopback('0.0.0.0'))
        self.assertFalse(_is_loopback('192.168.1.10'))

    def test_attempt_summary_never_exposes_command(self):
        summary = PipelineReadModel._attempt({
            'etapa': 'video',
            'status': 'submetida',
            'modelo': 'veo-fast',
            'comando': ['gflow', '--project', 'segredo'],
        })

        self.assertEqual(summary['etapa'], 'video')
        self.assertNotIn('comando', summary)

    def test_latest_image_review_exposes_only_review_summary(self):
        summary = PipelineReadModel._latest_image_review([
            {'etapa': 'imagem', 'resultado': 'concluida'},
            {
                'etapa': 'revisar-imagem',
                'resultado': 'rejeitada',
                'justificativa': 'Produto deformado.',
                'imagem_sha256': 'a' * 64,
                'comando': ['não expor'],
            },
        ])

        self.assertEqual(summary['resultado'], 'rejeitada')
        self.assertEqual(summary['justificativa'], 'Produto deformado.')
        self.assertNotIn('comando', summary)

    def test_next_action_requires_sheet_and_bound_approval(self):
        clip = {
            'plan': {
                'pipeline': {
                    'aprovacao_necessaria': True,
                    'gerar_video': True,
                },
            },
        }

        action = PipelineReadModel._next_action(
            clip,
            {'imagem': {}},
            {'aprovacao': 'rejeitada', 'gerar_carrossel': ''},
            approval_bound=True,
        )

        self.assertEqual(action, 'revisar_e_aprovar_imagem')

    def test_next_action_prioritizes_authorized_local_carousel(self):
        clip = {
            'plan': {
                'pipeline': {
                    'aprovacao_necessaria': True,
                    'gerar_video': True,
                },
                'carrossel': {'ativo': True},
            },
        }

        action = PipelineReadModel._next_action(
            clip,
            {'imagem': {}},
            {'aprovacao': 'rejeitada', 'gerar_carrossel': 'sim'},
            approval_bound=False,
        )

        self.assertEqual(action, 'gerar_carrossel')

    def test_next_action_ignores_sheet_carousel_without_active_plan(self):
        clip = {
            'plan': {
                'pipeline': {
                    'aprovacao_necessaria': True,
                    'gerar_video': True,
                },
            },
        }

        action = PipelineReadModel._next_action(
            clip,
            {'imagem': {}},
            {'aprovacao': 'rejeitada', 'gerar_carrossel': 'sim'},
            approval_bound=False,
        )

        self.assertEqual(action, 'revisar_e_aprovar_imagem')

    def test_video_always_requires_bound_approval_even_if_plan_flag_is_false(self):
        clip = {
            'plan': {
                'pipeline': {
                    'aprovacao_necessaria': False,
                    'gerar_video': True,
                },
            },
        }

        action = PipelineReadModel._next_action(
            clip,
            {'imagem': {}},
            {'aprovacao': 'aprovada', 'gerar_carrossel': ''},
            approval_bound=False,
        )

        self.assertEqual(action, 'revisar_e_aprovar_imagem')

    def test_legacy_invalid_state_is_reported_without_hiding_row(self):
        book = load_workbook(self.sheet)
        sheet = book['Controle']
        sheet.cell(4, HEADERS.index('video_status') + 1, 'aguardando')
        book.save(self.sheet)

        view = PipelineReadModel(self.config).productions()

        self.assertEqual(view['planilha']['linhas'], 1)
        self.assertEqual(len(view['erros']), 1)
        self.assertTrue(any(item['ativa'] for item in view['producoes']))

    def test_invalid_execution_state_recommends_correction(self):
        state_path = self.clip_dir / 'execucao.json'
        state_path.write_text(
            json.dumps({'schema_version': 2, 'fingerprint': 'd' * 64}),
            encoding='utf-8',
        )

        clip = PipelineReadModel(self.config).clips('PROD_01')[0]

        self.assertTrue(clip['erros'])
        self.assertEqual(clip['proxima_acao'], 'corrigir_estado_invalido')


if __name__ == '__main__':
    unittest.main()
