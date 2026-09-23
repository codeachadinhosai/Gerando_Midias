import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline_flow.services.gerar_modelo_planilha import OWNERS
from pipeline_flow.services.preparar_insumos import (
    HEADERS,
    Workbook,
    digest,
    update_spreadsheet_structure,
)


class SpreadsheetStructureUpdateTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'controle.xlsx'

        workbook = OpenpyxlWorkbook()
        control = workbook.active
        control.title = 'Controle'
        control.append(HEADERS)
        control.append(['ANTIGO'] * len(HEADERS))
        control.append(['Descricao antiga'] * len(HEADERS))
        values = [''] * len(HEADERS)
        values[HEADERS.index('classifica')] = 'sim'
        values[HEADERS.index('arquivo')] = 'produto.png'
        values[HEADERS.index('producao_id')] = 'PROD_01'
        values[HEADERS.index('ordem')] = 1
        control.append(values)
        control['B4'].hyperlink = 'https://example.com/produto.png'
        validation = DataValidation(type='list', formula1='"produto,inspiracao"')
        control.add_data_validation(validation)
        validation.add('I4:I500')
        control.freeze_panes = 'E32'

        guide = workbook.create_sheet('Guia')
        guide.append(['CAMPO', 'PREENCHIDO POR', 'OBRIGATORIO?', 'REGRA'])
        guide.append([
            'tipo_referencia',
            'VOCE',
            'Nao',
            'produto, inspiracao, detalhe ou ambiente',
        ])

        references = workbook.create_sheet('Referencias_Julia')
        references.append(['ARQUIVO', 'FUNCAO', 'GERENCIAMENTO'])
        references.append([
            'rosto.jpg',
            'Identidade facial principal',
            'GLOBAL / SISTEMA',
        ])

        example = workbook.create_sheet('Exemplo_Producao')
        example.append(['arquivo', 'produto_id'])
        example.append(['exemplo.jpg', 'EXEMPLO'])
        workbook.save(self.path)
        workbook.close()

    def test_update_preserves_operational_data_and_creates_backup(self):
        original_hash = digest(self.path)

        result = update_spreadsheet_structure(self.path)

        backup = Path(result['backup'])
        self.assertTrue(backup.is_file())
        self.assertEqual(digest(backup), original_hash)
        self.assertEqual(result['linhas_operacionais_preservadas'], 1)
        self.assertIn(
            'Referencias_Julia -> Referencias_Identidade',
            result['referencias_atualizadas'],
        )

        Workbook(self.path)
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertEqual(
            workbook.sheetnames,
            [
                'Controle',
                'Guia',
                'Referencias_Identidade',
                'Exemplo_Producao',
            ],
        )
        control = workbook['Controle']
        self.assertEqual(control['A4'].value, 'sim')
        self.assertEqual(control['B4'].value, 'produto.png')
        self.assertEqual(
            control['B4'].hyperlink.target,
            'https://example.com/produto.png',
        )
        self.assertEqual(control.freeze_panes, 'A4')
        self.assertEqual(
            [
                (selection.pane, selection.activeCell, str(selection.sqref))
                for selection in control.sheet_view.selection
            ],
            [('bottomLeft', 'A4', 'A4')],
        )
        self.assertEqual(control.auto_filter.ref, 'A1:AE500')
        self.assertEqual(control['Z2'].value, OWNERS['cta_destino'])
        self.assertEqual(control['AA2'].value, OWNERS['cta_palavra'])

        validations = {
            str(item.sqref): item.formula1
            for item in control.data_validations.dataValidation
        }
        self.assertIn('automatico', validations['F4:F500'])
        self.assertIn('base_edicao', validations['I4:I500'])
        self.assertIn('inspiração', validations['I4:I500'])
        self.assertEqual(len(control.conditional_formatting), 1)

        guide_values = {
            row[0]: row[2]
            for row in workbook['Guia'].iter_rows(
                min_row=2,
                values_only=True,
            )
        }
        self.assertIn('base_edicao', guide_values['tipo_referencia'])
        self.assertIn('inspiracao', guide_values['tipo_referencia'])

        references = workbook['Referencias_Identidade']
        self.assertEqual(references['A2'].value, 'identidade_rosto')
        self.assertEqual(references['B2'].value, 'rosto.jpg')
        self.assertEqual(
            workbook['Exemplo_Producao'].max_column,
            len(HEADERS),
        )
        workbook.close()

        updated_hash = digest(self.path)
        repeated = update_spreadsheet_structure(self.path)
        self.assertEqual(repeated['resultado'], 'planilha ja sincronizada')
        self.assertIsNone(repeated['backup'])
        self.assertEqual(digest(self.path), updated_hash)


if __name__ == '__main__':
    unittest.main()
