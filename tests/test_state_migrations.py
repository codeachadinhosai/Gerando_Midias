import argparse
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook as OpenpyxlWorkbook

from pipeline_flow.services import executar_flow as flow
from pipeline_flow.services import migrar_estados as migrations
from pipeline_flow.services.preparar_insumos import (
    HEADERS,
    Invalid,
    Workbook,
)


def create_workbook(path, clip, approval=''):
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = 'Controle'
    sheet.append(HEADERS)
    sheet.append([''] * len(HEADERS))
    sheet.append([''] * len(HEADERS))
    values = {
        'classifica': 'sim',
        'producao_id': clip['plan']['producao_id'],
        'ordem': str(clip['plan']['ordem']),
        'plano_arquivo': str(
            clip['folder'] / 'plano_clipe.json'
        ),
        'aprovacao': approval,
    }
    sheet.append([values.get(name, '') for name in HEADERS])
    book.save(path)


class StateMigrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name) / 'CLIP_01'
        self.folder.mkdir()
        (self.folder / 'plano_clipe.json').write_text(
            '{}',
            encoding='utf-8',
        )
        self.frame = self.folder / 'frame.png'
        self.frame.write_bytes(b'legacy frame')
        self.clip = {
            'folder': self.folder,
            'fingerprint': 'a' * 64,
            'revision_sha256': 'c' * 64,
            'flow': {'pacote_sha256': 'b' * 64},
            'plan': {
                'producao_id': 'PROD_01',
                'id_clipe': 'CLIP_01',
                'ordem': 1,
                'pipeline': {},
            },
        }
        self.state_path = self.folder / 'execucao.json'
        self.legacy_state = {
            'fingerprint': self.clip['fingerprint'],
            'imagem': {
                'arquivo': str(self.frame),
                'sha256': flow.digest(self.frame),
                'modelo': 'nano2',
            },
            'historico': [],
        }
        flow.write_json(self.state_path, self.legacy_state)

    def test_dry_run_reports_without_writing(self):
        original = self.state_path.read_bytes()

        result = migrations.migrate_clip_state(
            self.clip,
            apply=False,
        )

        self.assertEqual(result['estado'], 'pendente')
        self.assertEqual(
            result['alteracoes'],
            ['estado_v2', 'imagem_v2'],
        )
        self.assertEqual(self.state_path.read_bytes(), original)
        self.assertFalse(
            (self.folder / 'logs' / 'migrations').exists()
        )

    def test_apply_backs_up_and_is_idempotent(self):
        original_hash = flow.digest(self.state_path)

        first = migrations.migrate_clip_state(
            self.clip,
            apply=True,
        )
        migrated = flow.read_json(self.state_path)

        self.assertTrue(first['aplicada'])
        self.assertEqual(
            migrated['schema_version'],
            flow.EXECUTION_SCHEMA_VERSION,
        )
        self.assertEqual(
            migrated['imagem']['schema_version'],
            flow.MEDIA_SCHEMA_VERSION,
        )
        self.assertEqual(
            migrated['imagem']['revisao_sha256'],
            self.clip['revision_sha256'],
        )
        self.assertEqual(
            flow.digest(Path(first['backup'])),
            original_hash,
        )
        self.assertFalse((self.folder / '.execucao.lock').exists())

        second = migrations.migrate_clip_state(
            self.clip,
            apply=True,
        )
        repeated = flow.read_json(self.state_path)

        self.assertFalse(second['aplicada'])
        self.assertEqual(second['estado'], 'atual')
        self.assertEqual(len(repeated['historico_migracoes']), 1)
        self.assertEqual(
            len(list((self.folder / 'logs' / 'migrations').iterdir())),
            1,
        )

    def test_legacy_approval_is_archived_and_cleared_in_workbook(self):
        state = copy.deepcopy(self.legacy_state)
        state['aprovacao'] = {
            'arquivo': str(self.frame),
            'sha256': flow.digest(self.frame),
            'aprovada_em': '2026-01-01T00:00:00Z',
        }
        flow.write_json(self.state_path, state)
        sheet = self.folder / 'controle.xlsx'
        create_workbook(sheet, self.clip, approval='aprovada')

        result = migrations.migrate_clip_state(
            self.clip,
            sheet,
            apply=True,
        )

        migrated = flow.read_json(self.state_path)
        row = Workbook(sheet).records()[0]
        self.assertTrue(result['reaprovacao_necessaria'])
        self.assertTrue(result['aprovacao_planilha_limpa'])
        self.assertNotIn('aprovacao', migrated)
        self.assertEqual(
            migrated['aprovacoes_historicas'][-1]['resultado'],
            'revogada_por_migracao_v2',
        )
        self.assertEqual(row['aprovacao'], '')

    def test_system_cannot_grant_approval(self):
        sheet = self.folder / 'controle.xlsx'
        create_workbook(sheet, self.clip)
        workbook = Workbook(sheet)

        with self.assertRaisesRegex(Invalid, 'nunca'):
            workbook.set(4, 'aprovacao', 'aprovada')

    def test_explicit_human_path_accepts_only_review_decisions(self):
        sheet = self.folder / 'controle.xlsx'
        create_workbook(sheet, self.clip)
        workbook = Workbook(sheet)
        workbook.set_human_approval(4, 'rejeitada')
        workbook.save()

        self.assertEqual(Workbook(sheet).records()[0]['aprovacao'], 'rejeitada')
        with self.assertRaisesRegex(Invalid, 'aprovada ou rejeitada'):
            Workbook(sheet).set_human_approval(4, 'pendente')

    def test_corrupt_media_blocks_without_writing_backup(self):
        original = self.state_path.read_bytes()
        self.frame.write_bytes(b'tampered')

        with self.assertRaisesRegex(Invalid, 'Midia ausente ou alterada'):
            migrations.migrate_clip_state(
                self.clip,
                apply=True,
            )

        self.assertEqual(self.state_path.read_bytes(), original)
        self.assertFalse(
            (self.folder / 'logs' / 'migrations').exists()
        )

    def test_unknown_execution_schema_fails_closed(self):
        state = copy.deepcopy(self.legacy_state)
        state['schema_version'] = 999
        flow.write_json(self.state_path, state)

        with self.assertRaisesRegex(Invalid, 'Versao do estado'):
            flow.state_for(self.clip)

    def test_active_execution_lock_blocks_migration(self):
        handle = flow.acquire_clip_lock(self.clip, 'imagem')
        try:
            with self.assertRaisesRegex(Invalid, 'processo ativo'):
                migrations.migrate_clip_state(
                    self.clip,
                    apply=True,
                )
            self.assertEqual(
                flow.read_json(self.state_path),
                self.legacy_state,
            )
        finally:
            flow.release_lock(handle)


class RecoveryIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.folder = self.root / 'CLIP_01'
        self.folder.mkdir()
        (self.folder / 'plano_clipe.json').write_text(
            '{}',
            encoding='utf-8',
        )
        self.clip = {
            'folder': self.folder,
            'fingerprint': 'a' * 64,
            'revision_sha256': 'c' * 64,
            'flow': {
                'pacote_sha256': 'b' * 64,
                'referencias': [],
            },
            'refs': [],
            'prompts': {'imagem': 'prompt local'},
            'plan': {
                'producao_id': 'PROD_01',
                'id_clipe': 'CLIP_01',
                'ordem': 1,
                'pipeline': {
                    'gerar_imagem': True,
                    'metodo_imagem': 'i2i',
                    'gerar_video': False,
                    'metodo_video': None,
                    'aprovacao_necessaria': True,
                },
            },
        }
        self.sheet = self.root / 'controle.xlsx'
        create_workbook(self.sheet, self.clip, approval='aprovada')
        self.run_dir = self.folder / 'gerados' / 'interrupted'
        self.run_dir.mkdir(parents=True)
        self.output = self.run_dir / 'imagem.png'
        self.output.write_bytes(b'completed image')
        key = flow.operation_idempotency_key(
            self.clip,
            'imagem',
            'nano2',
        )
        flow.write_json(
            self.folder / 'execucao.json',
            {
                'schema_version': flow.EXECUTION_SCHEMA_VERSION,
                'fingerprint': self.clip['fingerprint'],
                'tentativa': {
                    'schema_version': flow.ATTEMPT_SCHEMA_VERSION,
                    'etapa': 'imagem',
                    'status': 'submetida',
                    'inicio': '2026-09-20T10:00:00Z',
                    'comando': ['gflow', 'image'],
                    'diretorio': str(self.run_dir),
                    'saida': str(self.output),
                    'modelo': 'nano2',
                    'idempotency_key': key,
                    'revisao_sha256': self.clip['revision_sha256'],
                    'fingerprint_sha256': self.clip['fingerprint'],
                },
            },
        )
        self.args = argparse.Namespace(
            acao='imagem',
            arquivo=None,
            planilha=self.sheet,
            gflow_raiz=self.root,
            projeto='unused',
            modelo_video='omni-flash',
            timeout=1,
            delivery_dir=self.root / 'entregas',
        )

    def test_submitted_output_recovers_through_real_workbook(self):
        with patch.object(flow.subprocess, 'run') as run:
            result = flow.execute(self.clip, self.args)

        run.assert_not_called()
        self.assertEqual(
            result,
            'tentativa recuperada sem nova chamada ao Flow',
        )
        state = flow.read_json(self.folder / 'execucao.json')
        row = Workbook(self.sheet).records()[0]
        self.assertNotIn('tentativa', state)
        self.assertEqual(
            state['historico'][-1]['resultado'],
            'recuperada_sem_reenvio',
        )
        self.assertEqual(row['imagem_status'], 'gerada')
        self.assertEqual(row['status'], 'aguardando_aprovacao')
        self.assertEqual(row['aprovacao'], '')
        self.assertTrue(
            (self.root / 'entregas' / 'indice_entregas.json').is_file()
        )
        self.assertFalse((self.folder / '.execucao.lock').exists())


if __name__ == '__main__':
    unittest.main()
